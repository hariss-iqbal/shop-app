#!/usr/bin/env python3
"""
stock_import.py — import the shop's Excel stock sheet into the local Supabase DB,
with GSMArena normalization (model names, marketing colors, RAM), or compare the
sheet against the current DB.

Subcommands:
  overview   Parse the Excel + read DB counts. Print what an import/compare would do.
             (No DB writes, no GSMArena needed.)
  prepare    Fetch GSMArena specs per model, build the generic->marketing color map,
             and write the normalized dataset. Flags UNCERTAIN color mappings.
             (Needs the GSMArena backend reachable at $STOCK_API. No DB writes.)
  compare    Compare the normalized Excel dataset against the current DB. (Run prepare first.)
  load       DESTRUCTIVE: wipe products+variants+models, then load the normalized dataset
             (slug auto-filled by DB trigger). Requires --yes. (Run prepare first.)

Config via env:
  STOCK_XLSX  path to the Excel sheet     (default: ~/Downloads/Stock Management.xlsx)
  STOCK_DB    postgres URL                (default: local supabase)
  STOCK_API   GSMArena backend base URL   (default: http://localhost:3001)
  STOCK_WORK  scratch dir for json/sql    (default: /tmp/stock-import)

The Excel is expected to have columns:
  Name, Type, Brand, Color, Storage (GB), RAM (GB), Condition, Condition Rating (1-10),
  Battery Health (%), PTA Status, IMEI / Serial, Quantity, Cost Price (PKR),
  Selling Price (PKR), Notes
"""
import os, sys, re, json, time, subprocess, urllib.request
from collections import defaultdict, OrderedDict

XLSX = os.environ.get("STOCK_XLSX", os.path.expanduser("~/Downloads/Stock Management.xlsx"))
DB   = os.environ.get("STOCK_DB", "postgresql://postgres:postgres@127.0.0.1:54322/postgres")
API  = os.environ.get("STOCK_API", "http://localhost:3001").rstrip("/")
WORK = os.environ.get("STOCK_WORK", "/tmp/stock-import")
os.makedirs(WORK, exist_ok=True)

# ----------------------------------------------------------------------------- helpers
def psql(sql):
    r = subprocess.run(["psql", DB, "-t", "-A", "-F", "\t", "-c", sql],
                       capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError("psql error: " + r.stderr.strip())
    return [ln.split("\t") for ln in r.stdout.strip().split("\n") if ln.strip()]

def psql_file(path):
    r = subprocess.run(["psql", DB, "-v", "ON_ERROR_STOP=1", "-f", path],
                       capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError("psql load error: " + r.stderr.strip())
    return r.stdout

def api_post(path, body):
    req = urllib.request.Request(API + path, data=json.dumps(body).encode(),
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.load(resp)

def api_up():
    try:
        urllib.request.urlopen(API + "/health", timeout=4); return True
    except Exception:
        return False

# ------------------------------------------------------------------- canonicalization
# Vocab is matched case-insensitively (sheets vary: "PTA"/"pta", "Box"/"open box").
PTA  = {"pta approved": "pta_approved", "pta": "pta_approved",
        "non-pta": "non_pta", "non pta": "non_pta",
        "online approved": "non_pta"}   # shop's "online approved" -> Non-PTA (per import choice)
COND = {"new": "new", "used": "used", "open box": "open_box", "box": "open_box",
        "sealed box": "new", "sealed": "new", "refurbished": "refurbished"}

def pta_norm(v):
    if v is None or str(v).strip() == "": return None
    return PTA.get(str(v).strip().lower(), str(v).strip())

def cond_norm(v):
    if v is None or str(v).strip() == "": return "new"   # blank condition -> new (column is NOT NULL)
    return COND.get(str(v).strip().lower(), str(v).strip().lower())

BRANDS_KNOWN = {"google": "Google", "apple": "Apple", "nothing": "Nothing",
                "oneplus": "OnePlus", "samsung": "Samsung", "xiaomi": "Xiaomi"}
def canon_brand(b):
    s = re.sub(r"\s+", " ", str(b or "").strip())
    return BRANDS_KNOWN.get(s.lower(), s)

ACCESSORY_KW = ("buds", "earbud", "watch", "case", "charger", "cable", "adapter")
def is_accessory(name):
    n = str(name or "").lower()
    return any(k in n for k in ACCESSORY_KW)

def canon_model(brand, name):
    """Light cleanup for grouping + a GSMArena search query. Final display name
    comes from GSMArena (gsm_name) during `prepare`."""
    s = re.sub(r"\s+", " ", str(name).strip().lower())
    bl = str(brand).strip().lower()
    if bl and s.startswith(bl + " "):
        s = s[len(bl):].strip()                        # drop a leading brand prefix in the name
    s = re.sub(r"\bcp\b", "", s).strip()              # shop convention: CP = Non-PTA marker
    s = re.sub(r"\bjv\b", "", s).strip()              # "JV" = regional (Japan) SKU marker, not the model
    s = re.sub(r"\b512\b", "", s).strip()              # redundant storage suffix in name
    s = re.sub(r"pixel\s+(\d+)\s+a\b", r"pixel \1a", s)  # "pixel 7 a" -> "pixel 7a"
    s = re.sub(r"^(oneplus)(\d.*)$", r"\1 \2", s)         # "oneplus13t" -> "oneplus 13t"
    s = re.sub(r"\s+", " ", s).strip()
    out = []
    for tok in s.split(" "):
        if tok in ("pro", "max", "plus", "ultra", "mini"): out.append(tok.capitalize())
        elif tok == "xl": out.append("XL")
        elif tok == "se": out.append("SE")
        elif tok == "iphone": out.append("iPhone")
        elif tok == "oneplus": out.append("OnePlus")
        elif tok == "pixel": out.append("Pixel")
        elif tok == str(brand).lower(): continue          # strip brand token from model name
        elif re.fullmatch(r"\d+a", tok): out.append(tok)  # a-series stays lowercase a
        elif re.fullmatch(r"\d+t", tok): out.append(tok[:-1] + "T")
        else: out.append(tok)
    return " ".join(out)

def canon_color(c):
    s = re.sub(r"\s+", " ", str(c).strip().lower()).replace("bule", "blue")
    return " ".join(w.capitalize() for w in s.split(" "))

def slugify(*parts):
    out = [re.sub(r"[^a-z0-9]+", "-", str(p).lower()) for p in parts if p not in (None, "")]
    return re.sub(r"-+", "-", "-".join(out)).strip("-")

# generic family -> candidate GSMArena marketing keywords (priority order)
COLOR_KW = {
    "black": ["obsidian", "stormy black", "midnight", "graphite", "charcoal", "matte space gray", "space gray", "black"],
    "white": ["snow", "porcelain", "cloudy white", "chalk", "starlight", "matte silver", "silver", "white"],
    "blue":  ["bay", "sierra blue", "pacific blue", "sea", "indigo", "blue"],
    "green": ["alpine green", "matte midnight green", "midnight green", "wintergreen", "sage", "mint", "seafoam", "green"],
    "gold":  ["matte gold", "gold"],
    "pink":  ["peony", "rose quartz", "rose", "kinda coral", "coral", "pink"],
    "purple": ["iris", "indigo", "purple"],
    "gray":  ["graphite", "space gray", "gray"],
    "silver": ["matte silver", "silver"],
    "hazel": ["hazel"], "moonstone": ["moonstone"], "jade": ["jade"],
    "lemon grass": ["lemongrass"], "lemongrass": ["lemongrass"],
}

def map_color(generic, gsm_colors, overrides, brand, model):
    if (brand, model, generic) in overrides:
        return overrides[(brand, model, generic)], "override"
    g = generic.lower().strip()
    norm = lambda s: re.sub(r"\s+", "", s.lower())
    for c in gsm_colors:
        if norm(c) == norm(g): return c, "exact"
    for kw in COLOR_KW.get(g, []):
        for c in gsm_colors:
            if kw in c.lower(): return c, "mapped"
    for c in gsm_colors:
        if g in c.lower(): return c, "loose"
    return None, "UNCERTAIN"

# -------------------------------------------------------------------------- parse Excel
COLS = ["Name", "Type", "Brand", "Color", "Storage (GB)", "RAM (GB)", "Condition",
        "Condition Rating (1-10)", "Battery Health (%)", "PTA Status", "IMEI / Serial",
        "Quantity", "Cost Price (PKR)", "Selling Price (PKR)", "Notes"]

def parse_excel():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.worksheets[0]
    ncol = ws.max_column                               # honor the sheet's actual column count
    hdr = [ws.cell(row=1, column=c).value for c in range(1, ncol + 1)]
    rows, skipped_acc = [], 0
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ncol + 1)]
        if all(v is None for v in vals):
            continue
        d = {k: (v.strip() if isinstance(v, str) else v) for k, v in zip(hdr, vals)}
        if is_accessory(d.get("Name")):                # non-phone items (e.g. earbuds) excluded
            skipped_acc += 1
            continue
        rows.append((r, d))
    if skipped_acc:
        print(f"(skipped {skipped_acc} non-phone/accessory row(s))", file=sys.stderr)

    def vkey(d):
        brand = canon_brand(d.get("Brand"))
        return (brand, canon_model(brand, d["Name"]),
                int(d["Storage (GB)"]) if d.get("Storage (GB)") is not None else None,
                pta_norm(d.get("PTA Status")),
                cond_norm(d.get("Condition")))

    groups = defaultdict(list)
    for r, d in rows:
        groups[vkey(d)].append((r, d))

    # price conflicts: same variant key, >1 distinct selling price -> set aside
    conflicts = {k for k, items in groups.items()
                 if len({i[1]["Selling Price (PKR)"] for i in items}) > 1}

    variants = OrderedDict()
    for k, items in groups.items():
        if k in conflicts:
            continue
        brand, model, storage, pta, cond = k
        sp = float(items[0][1]["Selling Price (PKR)"])
        variants[k] = {
            "brand": brand, "model": model, "storage_gb": storage,
            "pta_status": pta, "condition": cond, "selling_price": sp,
            "rows": [{
                "excel_row": r,
                "color": canon_color(d["Color"]),
                "ram_gb": int(d["RAM (GB)"]) if d.get("RAM (GB)") is not None else None,
                "cost_price": float(d["Cost Price (PKR)"]) if d.get("Cost Price (PKR)") is not None else 0.0,
                "condition_rating": int(d["Condition Rating (1-10)"]) if d.get("Condition Rating (1-10)") is not None else None,
                "battery_health": int(d["Battery Health (%)"]) if d.get("Battery Health (%)") is not None else None,
                "imei": (str(d["IMEI / Serial"]).strip() if d.get("IMEI / Serial") is not None else None),
                "qty": int(d["Quantity"]) if d.get("Quantity") is not None else 1,
                "notes": (str(d["Notes"]).strip() if d.get("Notes") is not None else None),
            } for r, d in items],
        }
    conflict_detail = [{"key": k, "rows": [r for r, _ in groups[k]],
                        "prices": sorted({i[1]["Selling Price (PKR)"] for i in groups[k]})}
                       for k in conflicts]
    return variants, conflict_detail

def db_counts():
    try:
        rows = psql("SELECT (SELECT count(*) FROM models),(SELECT count(*) FROM variants),"
                    "(SELECT count(*) FROM products),(SELECT count(*) FROM products WHERE variant_id IS NULL)")
        m, v, p, o = rows[0]
        return dict(models=int(m), variants=int(v), products=int(p), orphans=int(o))
    except Exception as e:
        return {"error": str(e)}

# ------------------------------------------------------------------------- subcommands
def cmd_overview():
    variants, conflicts = parse_excel()
    units = sum(r["qty"] for v in variants.values() for r in v["rows"])
    brands = sorted({v["brand"] for v in variants.values()})
    models = sorted({(v["brand"], v["model"]) for v in variants.values()})
    db = db_counts()
    print("=" * 64)
    print("STOCK IMPORT — OVERVIEW")
    print("=" * 64)
    print(f"Excel file : {XLSX}")
    print(f"\nExcel (clean, loadable):")
    print(f"  brands   : {', '.join(brands)}")
    print(f"  models   : {len(models)}")
    print(f"  variants : {len(variants)}   (grouped by model+storage+PTA+condition)")
    print(f"  units    : {units}")
    if conflicts:
        print(f"\n  ⚠ price-conflict variant groups SKIPPED ({len(conflicts)}):")
        for c in conflicts:
            b, m, st, pta, cond = c["key"]
            print(f"     - {m} {st}GB {cond}/{pta}  excel rows {c['rows']}  prices {c['prices']}")
    print(f"\nCurrent local DB:")
    if "error" in db:
        print(f"  (could not read DB: {db['error']})")
    else:
        print(f"  models={db['models']}  variants={db['variants']}  products={db['products']}  orphans={db['orphans']}")
    print(f"\nWhat each option does:")
    print(f"  • IMPORT ALL : wipe DB models/variants/products, GSMArena-normalize, load {len(variants)} variants / {units} units.")
    print(f"  • COMPARE    : GSMArena-normalize the sheet and diff against the current DB (no writes).")
    print(f"\nGSMArena backend ($STOCK_API={API}): {'reachable' if api_up() else 'NOT reachable — needed for import/compare'}")

def cmd_prepare():
    if not api_up():
        print(f"ERROR: GSMArena backend not reachable at {API}. Start the shop backend "
              f"(cd backend && npm run dev) or set STOCK_API.", file=sys.stderr)
        sys.exit(2)
    variants, conflicts = parse_excel()
    models = sorted({(v["brand"], v["model"]) for v in variants.values()})

    # load optional human overrides for color mapping
    ov_path = os.path.join(WORK, "color_overrides.json")
    overrides = {}
    if os.path.exists(ov_path):
        for k, val in json.load(open(ov_path)).items():
            b, m, g = k.split("|"); overrides[(b, m, g)] = val

    specs = {}
    print(f"Fetching GSMArena specs for {len(models)} models...")
    for b, m in models:
        try:
            res = api_post("/api/products/fetch-specs", {"brand": b, "model": m})
            d = res.get("data") or {}
            specs[(b, m)] = {"gsm_name": d.get("modelName") or m, "colors": d.get("colors", []),
                             "storage": d.get("storage", []), "ram": d.get("ram", [])}
            print(f"  ok  {b} {m} -> {d.get('modelName')!r} colors={d.get('colors')}")
        except Exception as e:
            specs[(b, m)] = {"gsm_name": m, "colors": [], "storage": [], "ram": []}
            print(f"  !!  {b} {m} -> fetch failed ({e}); using raw name, no color mapping")
        time.sleep(0.4)  # GSMArena throttles rapid calls

    norm = {"models": [], "variants": []}
    color_review, uncertain = [], 0
    for v in variants.values():
        sp = specs[(v["brand"], v["model"])]
        gsm_name, gsm_ram, gsm_storage = sp["gsm_name"], sp["ram"], sp["storage"]
        def pick_ram(excel_ram):
            if len(gsm_ram) == 1: return gsm_ram[0]
            if len(gsm_ram) == 0: return excel_ram
            return excel_ram if excel_ram in gsm_ram else (gsm_ram[-1] if excel_ram is None else excel_ram)
        nv = {**{k: v[k] for k in ("brand", "storage_gb", "pta_status", "condition", "selling_price")},
              "model": gsm_name, "orig_model": v["model"],
              "predicted_slug": slugify(v["brand"], gsm_name,
                                        (str(v["storage_gb"]) + "gb") if v["storage_gb"] else None,
                                        v["condition"], v["pta_status"]),
              "rows": []}
        for r in v["rows"]:
            mk, conf = map_color(r["color"], sp["colors"], overrides, v["brand"], v["model"])
            if conf in ("UNCERTAIN", "loose"): uncertain += 1
            color_review.append({"brand": v["brand"], "model": gsm_name, "generic": r["color"],
                                 "marketing": mk, "confidence": conf, "gsm_colors": sp["colors"]})
            nv["rows"].append({**r, "generic_color": r["color"], "color": mk or r["color"],
                               "ram_gb": pick_ram(r["ram_gb"])})
        norm["variants"].append(nv)
    norm["models"] = sorted({(v["brand"], v["model"]) for v in norm["variants"]})
    json.dump(norm, open(os.path.join(WORK, "dataset_norm.json"), "w"), indent=2, default=str)
    json.dump(color_review, open(os.path.join(WORK, "color_review.json"), "w"), indent=2)

    print(f"\nNormalized -> {WORK}/dataset_norm.json  "
          f"({len(norm['models'])} models, {len(norm['variants'])} variants)")
    print("\nCOLOR MAP (generic -> GSMArena marketing):")
    for c in sorted(color_review, key=lambda x: (x["confidence"] != "UNCERTAIN", x["model"], x["generic"])):
        flag = "   <== REVIEW" if c["confidence"] in ("UNCERTAIN", "loose") else ""
        print(f"  {c['brand']:7} {c['model']:16} {c['generic']:12} -> {str(c['marketing']):16} "
              f"[{c['confidence']}]{flag}")
    if uncertain:
        print(f"\n⚠ {uncertain} mapping(s) need review. To override, write {ov_path} as "
              f'{{"Brand|Model|GenericColor": "MarketingColor", ...}} and re-run prepare.')

def _snapshot():
    out = psql("""
    SELECT b.name, m.name, v.storage_gb, v.pta_status, v.condition::text, v.selling_price,
           v.slug, v.stock_count,
           coalesce((SELECT string_agg(c,'|' ORDER BY c) FROM unnest(v.available_colors) c),'')
    FROM variants v JOIN models m ON v.model_id=m.id JOIN brands b ON m.brand_id=b.id
    """)
    snap = {}
    for r in out:
        storage = int(r[2]) if r[2] not in (None, "") else None
        snap[(r[0], r[1], storage, r[3], r[4])] = dict(
            sell=float(r[5]), slug=r[6], stock=int(r[7]), colors=r[8])
    return snap

def cmd_compare():
    path = os.path.join(WORK, "dataset_norm.json")
    if not os.path.exists(path):
        print("ERROR: run `prepare` first (need dataset_norm.json).", file=sys.stderr); sys.exit(2)
    norm = json.load(open(path))
    excel = {}
    for v in norm["variants"]:
        excel[(v["brand"], v["model"], v["storage_gb"], v["pta_status"], v["condition"])] = {
            "sell": float(v["selling_price"]),
            "units": sum(r["qty"] for r in v["rows"]),
            "colors": "|".join(sorted({r["color"] for r in v["rows"]})),
            "slug": v["predicted_slug"]}
    db = _snapshot()
    only_excel = sorted(set(excel) - set(db))
    only_db = sorted(set(db) - set(excel))
    both = sorted(set(excel) & set(db))
    print("=" * 64); print("COMPARE — normalized Excel vs current DB"); print("=" * 64)
    print(f"\nIn Excel but NOT in DB (would be ADDED) — {len(only_excel)}:")
    for k in only_excel: print(f"  + {k[1]} {k[2]}GB {k[4]}/{k[3]}  ({excel[k]['units']} units, {excel[k]['colors']})")
    print(f"\nIn DB but NOT in Excel (extra in DB) — {len(only_db)}:")
    for k in only_db: print(f"  - {k[1]} {k[2]}GB {k[4]}/{k[3]}  (db stock {db[k]['stock']})")
    print(f"\nIn BOTH — {len(both)} (field differences):")
    diffs = 0
    for k in both:
        e, d = excel[k], db[k]
        msg = []
        if abs(e["sell"] - d["sell"]) > 0.001: msg.append(f"sell {e['sell']:.0f}!=db {d['sell']:.0f}")
        if e["slug"] != d["slug"]: msg.append(f"slug {e['slug']}!=db {d['slug']}")
        if e["colors"] != d["colors"]: msg.append(f"colors [{e['colors']}]!=db[{d['colors']}]")
        if e["units"] != d["stock"]: msg.append(f"excel units {e['units']}!=db stock {d['stock']}")
        if msg: diffs += 1; print(f"  ~ {k[1]} {k[2]}GB {k[4]}/{k[3]}: {'; '.join(msg)}")
    if not diffs: print("  (all matching variants identical on sell/slug/colors/units)")

def _gen_sql(norm):
    def q(s): return "NULL" if s is None else "'" + str(s).replace("'", "''") + "'"
    def n(x): return "NULL" if x is None else str(x)
    L = ["BEGIN;"]
    brands = sorted({v["brand"] for v in norm["variants"]})
    L.append("INSERT INTO brands (name) SELECT v.name FROM (VALUES %s) AS v(name) "
             "WHERE NOT EXISTS (SELECT 1 FROM brands b WHERE b.name=v.name);"
             % ",".join("(%s)" % q(b) for b in brands))
    L.append("INSERT INTO models (brand_id,name) SELECT b.id,v.name FROM (VALUES %s) AS v(brand,name) "
             "JOIN brands b ON b.name=v.brand;"
             % ",".join("(%s,%s)" % (q(b), q(m)) for b, m in norm["models"]))
    vv = ",".join("(%s,%s,%s,%s,%s,%s)" % (q(v["brand"]), q(v["model"]), n(v["storage_gb"]),
                  q(v["pta_status"]), q(v["condition"]), n(v["selling_price"])) for v in norm["variants"])
    L.append("INSERT INTO variants (model_id,storage_gb,pta_status,condition,selling_price,is_active) "
             "SELECT m.id,v.storage_gb::integer,v.pta,v.cond::product_condition,v.sell::numeric,false "
             "FROM (VALUES %s) AS v(brand,model,storage_gb,pta,cond,sell) "
             "JOIN brands b ON b.name=v.brand JOIN models m ON m.brand_id=b.id AND m.name=v.model;" % vv)
    pr = []
    for v in norm["variants"]:
        for r in v["rows"]:
            pr.append("(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)" % (
                q(v["brand"]), q(v["model"]), n(v["storage_gb"]), q(v["pta_status"]), q(v["condition"]),
                q(r["color"]), n(r["ram_gb"]), n(r["cost_price"]), n(r.get("condition_rating")),
                n(r.get("battery_health")), n(r["qty"]), q(r.get("notes"))))
    L.append("INSERT INTO products (brand_id,model,model_id,variant_id,storage_gb,ram_gb,color,condition,"
             "pta_status,cost_price,selling_price,condition_rating,battery_health,status,product_type,notes) "
             "SELECT b.id,m.name,m.id,vr.id,r.storage_gb::integer,r.ram_gb::integer,r.color,r.cond::product_condition,"
             "r.pta::pta_status,r.cost::numeric,vr.selling_price,r.rating::integer,r.batt::integer,'available','phone',r.notes "
             "FROM (VALUES %s) AS r(brand,model,storage_gb,pta,cond,color,ram_gb,cost,rating,batt,qty,notes) "
             "JOIN brands b ON b.name=r.brand JOIN models m ON m.brand_id=b.id AND m.name=r.model "
             "JOIN variants vr ON vr.model_id=m.id AND vr.storage_gb IS NOT DISTINCT FROM r.storage_gb::integer "
             "AND vr.pta_status=r.pta AND vr.condition=r.cond::product_condition "
             "CROSS JOIN LATERAL generate_series(1,r.qty::integer) g;" % ",".join(pr))
    L.append("COMMIT;")
    return "\n".join(L)

def cmd_load(yes):
    path = os.path.join(WORK, "dataset_norm.json")
    if not os.path.exists(path):
        print("ERROR: run `prepare` first.", file=sys.stderr); sys.exit(2)
    if not yes:
        print("REFUSING: `load` is destructive (wipes models/variants/products). Pass --yes.", file=sys.stderr); sys.exit(2)
    norm = json.load(open(path))
    sql_path = os.path.join(WORK, "load.sql")
    open(sql_path, "w").write(_gen_sql(norm))
    print("Wiping models/variants/products and loading normalized data...")
    psql("DELETE FROM products; DELETE FROM variants; DELETE FROM models;")
    psql_file(sql_path)
    db = db_counts()
    units = sum(r["qty"] for v in norm["variants"] for r in v["rows"])
    print(f"Loaded. DB now: models={db['models']} variants={db['variants']} products={db['products']} "
          f"orphans={db['orphans']} (expected variants={len(norm['variants'])}, units={units})")
    noslug = psql("SELECT count(*) FROM variants WHERE slug IS NULL")[0][0]
    print(f"variants without slug: {noslug} (should be 0 — slugs auto-generated by DB trigger)")

# ------------------------------------------------------------------------------- main
def main():
    args = sys.argv[1:]
    cmd = args[0] if args else "overview"
    if cmd == "overview": cmd_overview()
    elif cmd == "prepare": cmd_prepare()
    elif cmd == "compare": cmd_compare()
    elif cmd == "load": cmd_load("--yes" in args)
    else:
        print(__doc__); sys.exit(1)

if __name__ == "__main__":
    main()
